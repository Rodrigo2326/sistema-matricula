# -*- coding: utf-8 -*-
"""
Genera el informe del avance en formato Word (.docx) a partir de su Markdown.

El Markdown es la unica fuente de verdad: este script solo lo maqueta, igual que
generar_informe.py hace con el PDF. La version en Word existe para poder
insertar las capturas de pantalla y los ajustes finales antes de entregar.

Uso:
    python docs/generar_word.py [apf1|apf2]

Sin argumento genera el informe del APF2, que es el avance vigente.

Requiere: python-docx (pip install python-docx)
"""

import io
import os
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

AQUI = os.path.dirname(os.path.abspath(__file__))

AZUL = RGBColor(0x1A, 0x2A, 0x52)
AZUL_CLARO = RGBColor(0x2E, 0x4A, 0x9E)
GRIS = RGBColor(0x55, 0x55, 0x55)

AVANCES = {
    'apf1': {
        'origen': 'informe-apf1.md',
        'destino': 'Informe_APF1.docx',
        'portada': 'Avance de Proyecto Final 1',
        'semana': '6',
        'version': None,
    },
    'apf2': {
        'origen': 'informe-apf2.md',
        'destino': 'Informe_APF2.docx',
        'portada': 'Avance de Proyecto Final 2',
        'semana': '11',
        'version': 'v2.0.0',
    },
}

INTEGRANTES = [
    'Ccahuana Huillca Giancarlo',
    'Aparcana Mamani Rodrigo Alonso',
    'Arnold Jhuncor Díaz Silva',
]


# --------------------------------------------------------------------------
# Utilidades de formato
# --------------------------------------------------------------------------

def preparar_estilos(documento):
    """Fija la tipografia base y el interlineado del documento."""
    normal = documento.styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.15

    for nivel, tam in (('Heading 1', 15), ('Heading 2', 12.5), ('Heading 3', 11.5)):
        estilo = documento.styles[nivel]
        estilo.font.name = 'Calibri'
        estilo.font.size = Pt(tam)
        estilo.font.bold = True
        estilo.font.color.rgb = AZUL
        estilo.paragraph_format.space_before = Pt(14)
        estilo.paragraph_format.space_after = Pt(6)


def escribir_con_marcas(parrafo, texto):
    """
    Vuelca el texto en el parrafo respetando las marcas de Markdown en linea:
    **negrita**, *cursiva* y `codigo`.
    """
    for trozo in re.split(r'(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)', texto):
        if not trozo:
            continue
        if trozo.startswith('**') and trozo.endswith('**'):
            parrafo.add_run(trozo[2:-2]).bold = True
        elif trozo.startswith('`') and trozo.endswith('`'):
            corrida = parrafo.add_run(trozo[1:-1])
            corrida.font.name = 'Consolas'
            corrida.font.size = Pt(9.5)
            corrida.font.color.rgb = AZUL_CLARO
        elif trozo.startswith('*') and trozo.endswith('*'):
            parrafo.add_run(trozo[1:-1]).italic = True
        else:
            parrafo.add_run(trozo)


def sombrear(celda, color):
    """Pinta el fondo de una celda; python-docx no lo expone directamente."""
    sombra = celda._tc.get_or_add_tcPr().makeelement(qn('w:shd'), {})
    sombra.set(qn('w:val'), 'clear')
    sombra.set(qn('w:fill'), color)
    celda._tc.get_or_add_tcPr().append(sombra)


# --------------------------------------------------------------------------
# Bloques
# --------------------------------------------------------------------------

def portada(documento, avance):
    titulo = documento.add_paragraph()
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    corrida = titulo.add_run(avance['portada'])
    corrida.bold = True
    corrida.font.size = Pt(26)
    corrida.font.color.rgb = AZUL

    subtitulo = documento.add_paragraph()
    subtitulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    corrida = subtitulo.add_run('Sistema de Matrícula')
    corrida.font.size = Pt(16)
    corrida.font.color.rgb = AZUL_CLARO

    documento.add_paragraph()

    datos = [
        ('Curso', 'Herramientas de Desarrollo (100000S66T)'),
        ('Periodo', '2026-II'),
        ('Semana', avance['semana']),
    ]
    if avance['version']:
        datos.append(('Versión', avance['version']))
    datos += [
        ('Proyecto', 'Sistema de Matrícula'),
        ('Repositorio', 'github.com/Rodrigo2326/sistema-matricula'),
    ]

    for etiqueta, valor in datos:
        parrafo = documento.add_paragraph()
        parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        parrafo.paragraph_format.space_after = Pt(2)
        corrida = parrafo.add_run('%s: %s' % (etiqueta, valor))
        corrida.font.size = Pt(10.5)
        corrida.font.color.rgb = GRIS

    documento.add_paragraph()

    cabecera = documento.add_paragraph()
    cabecera.alignment = WD_ALIGN_PARAGRAPH.CENTER
    corrida = cabecera.add_run('Integrantes')
    corrida.bold = True
    corrida.font.color.rgb = AZUL

    for nombre in INTEGRANTES:
        parrafo = documento.add_paragraph()
        parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        parrafo.paragraph_format.space_after = Pt(2)
        parrafo.add_run(nombre).font.size = Pt(11)

    documento.add_page_break()


def bloque_codigo(documento, lineas):
    parrafo = documento.add_paragraph()
    parrafo.paragraph_format.left_indent = Cm(0.6)
    parrafo.paragraph_format.space_before = Pt(6)
    corrida = parrafo.add_run('\n'.join(lineas))
    corrida.font.name = 'Consolas'
    corrida.font.size = Pt(9)
    corrida.font.color.rgb = AZUL


def bloque_tabla(documento, filas):
    """Primera fila como encabezado; la segunda del Markdown es el separador."""
    cuerpo = [f for f in filas[1:] if not re.match(r'^\s*\|[\s:|-]+\|\s*$', f)]
    encabezados = [c.strip() for c in filas[0].strip().strip('|').split('|')]

    tabla = documento.add_table(rows=1, cols=len(encabezados))
    tabla.style = 'Table Grid'
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER

    for celda, texto in zip(tabla.rows[0].cells, encabezados):
        celda.text = ''
        escribir_con_marcas(celda.paragraphs[0], texto)
        for parrafo in celda.paragraphs:
            for corrida in parrafo.runs:
                corrida.bold = True
                corrida.font.size = Pt(9.5)
                corrida.font.color.rgb = AZUL
        sombrear(celda, 'F4F6FB')

    for fila in cuerpo:
        valores = [c.strip() for c in fila.strip().strip('|').split('|')]
        celdas = tabla.add_row().cells
        for celda, texto in zip(celdas, valores):
            celda.text = ''
            escribir_con_marcas(celda.paragraphs[0], texto)
            for parrafo in celda.paragraphs:
                for corrida in parrafo.runs:
                    corrida.font.size = Pt(9.5)

    documento.add_paragraph().paragraph_format.space_after = Pt(4)


# --------------------------------------------------------------------------
# Conversion
# --------------------------------------------------------------------------

def convertir(documento, markdown):
    lineas = markdown.split('\n')
    i = 0
    # El encabezado del Markdown (titulo y metadatos) ya esta en la portada.
    while i < len(lineas) and lineas[i].strip() != '---':
        i += 1
    i += 1

    while i < len(lineas):
        linea = lineas[i]
        desnuda = linea.strip()

        if desnuda == '' or desnuda == '---':
            i += 1
            continue

        if desnuda.startswith('```'):
            i += 1
            codigo = []
            while i < len(lineas) and not lineas[i].strip().startswith('```'):
                codigo.append(lineas[i])
                i += 1
            i += 1
            bloque_codigo(documento, codigo)
            continue

        if desnuda.startswith('|'):
            filas = []
            while i < len(lineas) and lineas[i].strip().startswith('|'):
                filas.append(lineas[i])
                i += 1
            bloque_tabla(documento, filas)
            continue

        if desnuda.startswith('#'):
            nivel = len(desnuda) - len(desnuda.lstrip('#'))
            escribir_con_marcas(
                documento.add_heading('', min(nivel, 3)),
                desnuda.lstrip('#').strip())
            i += 1
            continue

        if desnuda.startswith('> '):
            cita = []
            while i < len(lineas) and lineas[i].strip().startswith('>'):
                cita.append(lineas[i].strip().lstrip('>').strip())
                i += 1
            parrafo = documento.add_paragraph()
            parrafo.paragraph_format.left_indent = Cm(0.8)
            escribir_con_marcas(parrafo, ' '.join(cita))
            for corrida in parrafo.runs:
                corrida.italic = True
                corrida.font.color.rgb = GRIS
            continue

        if re.match(r'^[-*] ', desnuda) or re.match(r'^\d+\. ', desnuda):
            numerada = bool(re.match(r'^\d+\. ', desnuda))
            while i < len(lineas):
                actual = lineas[i].strip()
                if not (re.match(r'^[-*] ', actual) or re.match(r'^\d+\. ', actual)):
                    break
                texto = re.sub(r'^([-*]|\d+\.)\s+', '', actual)
                i += 1
                # Continuaciones indentadas del mismo elemento.
                while i < len(lineas) and lineas[i].startswith('  ') and lineas[i].strip():
                    if re.match(r'^\s*([-*]|\d+\.)\s', lineas[i]):
                        break
                    texto += ' ' + lineas[i].strip()
                    i += 1
                estilo = 'List Number' if numerada else 'List Bullet'
                escribir_con_marcas(documento.add_paragraph(style=estilo), texto)
            continue

        # Parrafo normal: se unen las lineas hasta el proximo bloque.
        parrafo_texto = [desnuda]
        i += 1
        while i < len(lineas):
            siguiente = lineas[i].strip()
            if (siguiente == '' or siguiente.startswith(('#', '|', '>', '```', '---'))
                    or re.match(r'^([-*]|\d+\.)\s', siguiente)):
                break
            parrafo_texto.append(siguiente)
            i += 1
        parrafo = documento.add_paragraph()
        parrafo.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        escribir_con_marcas(parrafo, ' '.join(parrafo_texto))


def main():
    clave = (sys.argv[1].lower() if len(sys.argv) > 1 else 'apf2')
    if clave not in AVANCES:
        sys.exit('Avance desconocido: %s (use apf1 o apf2)' % clave)
    avance = AVANCES[clave]

    origen = os.path.join(AQUI, avance['origen'])
    destino = os.path.join(AQUI, avance['destino'])
    if not os.path.exists(origen):
        sys.exit('No se encontro el informe en %s' % origen)

    documento = Document()
    for seccion in documento.sections:
        seccion.left_margin = Cm(2.4)
        seccion.right_margin = Cm(2.4)
        seccion.top_margin = Cm(2.2)
        seccion.bottom_margin = Cm(2.2)

    preparar_estilos(documento)
    portada(documento, avance)
    convertir(documento, io.open(origen, encoding='utf-8').read())
    documento.save(destino)

    print('Word generado: %s' % destino)
    print('Parrafos: %d, tablas: %d' % (len(documento.paragraphs), len(documento.tables)))


if __name__ == '__main__':
    main()
